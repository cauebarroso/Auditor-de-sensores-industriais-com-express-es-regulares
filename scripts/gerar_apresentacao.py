"""
Gera a apresentação em dois formatos a partir da mesma descrição de slides:

  docs/Apresentacao.pptx  editável (PowerPoint / Google Slides), com o roteiro
                          de fala nas anotações de cada slide
  docs/Apresentacao.pdf   impressa pelo Microsoft Edge / Google Chrome

Requer: python-pptx, Pillow e markdown (pip install python-pptx pillow markdown).
Uso (na raiz do repositório):
  python scripts/gerar_apresentacao.py
"""

from __future__ import annotations

import html
import sys
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(RAIZ / "src"), str(RAIZ / "tests"), str(RAIZ / "scripts")]

from casos_teste import CASOS  # noqa: E402
from diagnostico import AUTOMATOS  # noqa: E402
from expressoes import POR_CODIGO  # noqa: E402
import auditor  # noqa: E402
from gerar_relatorio import capturar, resultados_pytest  # noqa: E402
from pdf import html_para_pdf  # noqa: E402

LARGURA, ALTURA = 13.333, 7.5
AZUL, AZUL_CLARO, CINZA, VERDE, VERMELHO = "1F3A5F", "EEF3F9", "5B6573", "1D6B3F", "B03A2E"
FONTE, FONTE_CODIGO = "Segoe UI", "Consolas"
DIAGRAMAS = RAIZ / "automatos" / "diagramas"


# ---------------------------------------------------------------------- #
# Descrição dos slides
# ---------------------------------------------------------------------- #
@dataclass
class Texto:
    x: float
    y: float
    w: float
    h: float
    linhas: list[str]          # "• " no início vira marcador; "**" no início = negrito
    tamanho: int = 18
    cor: str = "1B1F24"
    alinhamento: str = "esquerda"
    codigo: bool = False
    fundo: str | None = None


@dataclass
class Imagem:
    x: float
    y: float
    w: float
    h: float
    caminho: Path


@dataclass
class Tabela:
    x: float
    y: float
    w: float
    linhas: list[list[str]]    # a primeira linha é o cabeçalho
    larguras: list[float]
    tamanho: int = 13
    altura_linha: float = 0.36
    codigo_colunas: tuple[int, ...] = ()


@dataclass
class Slide:
    titulo: str
    elementos: list = field(default_factory=list)
    notas: str = ""
    capa: bool = False


def exemplos(codigo: str, n: int = 2) -> tuple[list[str], list[str]]:
    casos = CASOS[codigo]
    aceitas = [c.cadeia for c in casos if c.aceita][:n]
    rejeitadas = [f"{c.cadeia}  ← {c.motivo[0].lower()}{c.motivo[1:]}"
                  for c in casos if not c.aceita and c.cadeia.strip()][:n]
    return aceitas, rejeitadas


def slide_er(codigo: str, resumo: str, imagem: Path | None, notas: str, altura_formal: float = 1.25) -> Slide:
    """Slide de uma ER: resumo, ER formal, padrão do código, AFNε (se houver
    imagem) e exemplos de cadeias aceitas e rejeitadas."""
    er = POR_CODIGO[codigo]
    afn = AUTOMATOS[codigo]
    y_codigo = 1.85 + altura_formal + 0.45
    if imagem is None:
        # Sem diagrama: exemplos empilhados em largura total (cadeias do log são longas)
        y_exemplos = y_codigo + 1.15
        n = 4 if 6.95 - y_exemplos >= 1.9 else 3
        aceitas, rejeitadas = exemplos(codigo, n)
        corpo = [
            Texto(0.5, y_codigo + 0.75, 12.3, 0.35, ["Exemplos (tests/casos_teste.py)"], 13, AZUL),
            Texto(0.5, y_exemplos, 12.3, 0.22 * n, ["✔ " + a for a in aceitas], 13, VERDE, codigo=True),
            Texto(0.5, y_exemplos + 0.22 * n, 12.3, 0.22 * n, ["✘ " + r for r in rejeitadas], 13, VERMELHO,
                  codigo=True),
        ]
    else:
        aceitas, rejeitadas = exemplos(codigo)
        corpo = [
            Imagem(0.5, y_codigo + 0.65, 12.3, 6.65 - (y_codigo + 0.65), imagem),
            Texto(0.5, 6.7, 6.1, 0.6, ["✔ " + a for a in aceitas], 12, VERDE, codigo=True),
            Texto(6.7, 6.7, 6.1, 0.6, ["✘ " + r for r in rejeitadas], 12, VERMELHO, codigo=True),
        ]
    return Slide(
        f"{codigo} — {er.nome}",
        [
            Texto(0.5, 1.0, 12.3, 0.45, [resumo], 16, CINZA),
            Texto(0.5, 1.5, 1.6, 0.35, ["ER formal"], 13, AZUL),
            Texto(0.5, 1.85, 12.3, altura_formal, er.formal_completa(), 14, codigo=True, fundo=AZUL_CLARO),
            Texto(0.5, y_codigo - 0.35, 3.0, 0.35, ["No código (re.fullmatch)"], 13, AZUL),
            Texto(0.5, y_codigo, 12.3, 0.5, [f'r"{er.padrao}"'], 12 if len(er.padrao) > 110 else 14,
                  codigo=True, fundo=AZUL_CLARO),
            *corpo,
            Texto(10.3, 0.35, 2.6, 0.4, [f"AFNε: {afn.n_estados} estados · {afn.total_movimentos_vazios()} ε"],
                  12, "FFFFFF", "direita"),
        ],
        notas,
    )


def saida_demonstracao() -> list[str]:
    """Saída REAL do programa para os comandos mostrados no slide."""
    erro = capturar(auditor.main, ["--pacote", "CMD ATU-VLV-0201 AJUSTAR 120%"]).splitlines()
    erro = [linha for linha in erro if linha.strip() and "Pacote:" not in linha]
    traco = [linha for linha in capturar(auditor.explicar_afn, "ATU-VLV-0420", "ER-01").splitlines() if linha.strip()]
    inicio = next(i for i, linha in enumerate(traco) if "passo" in linha)
    passos = traco[inicio:inicio + 4] + ["      ⋮"] + [traco[inicio + 13], traco[inicio + 14].strip()]
    return [*erro, "", *passos]


def slide_testes() -> Slide:
    """Slide com os números reais da suíte (coletados do pytest)."""
    resumo, por_arquivo = resultados_pytest()
    descricao = {
        "tests/test_expressoes.py": "Cadeias aceitas/rejeitadas, requisitos 6+6, categoria única",
        "tests/test_automatos.py": "AFNε × re, provas de equivalência, .jff, linguagens disjuntas",
        "tests/test_aplicacao.py": "Entradas vazias/inválidas, diagnóstico, lote, CLI",
    }
    linhas = [["Arquivo", "O que verifica", "Testes"]]
    linhas += [[arquivo.split("/")[1], descricao.get(arquivo, ""), str(n)] for arquivo, n in sorted(por_arquivo.items())]
    linhas.append(["Total", "python -m pytest", f"{sum(por_arquivo.values())} ✔"])
    return Slide("Testes e resultados", [
        Tabela(0.5, 1.15, 12.35, linhas, [3.0, 7.75, 1.6], 16, 0.55, codigo_colunas=(0,)),
        Texto(0.5, 4.2, 6.0, 2.8, [
            "**Casos-limite cobertos",
            "• 0% e 100%; 850 e 1099 hPa",
            "• octetos 0 e 255; máscaras /0 e /32",
            "• 29/02, 30/04, 31/07 aceitas; 30/02, 31/04, hora 24 rejeitadas",
            "• menor e maior cadeia de L₁; a cadeia vazia ε",
        ], 17),
        Texto(6.85, 4.2, 6.0, 2.8, [
            "**Além das 80 cadeias de teste",
            "• JFLAP 7.1: 80/80 com o resultado esperado",
            "• 10.000 cadeias aleatórias: 0 divergências entre AFNε e re",
            "• Dados de exemplo: 20 íntegros e 14 corrompidos, todos com a causa correta",
        ], 17),
    ], f"Integrante 2 (~1 min). Resultado atual: {resumo}. Os casos de teste ficam em um único arquivo usado pelo "
       "pytest, pelo JFLAP e pelo relatório, então não há divergência entre eles.")


def slides() -> list[Slide]:
    return [
        Slide("", [
            Texto(0.8, 1.7, 11.7, 1.4, ["Auditor Léxico de Telemetria", "de Sensores Industriais"], 40, "FFFFFF"),
            Texto(0.8, 3.25, 11.7, 0.6, ["Um firewall léxico construído com Expressões Regulares e AFNε"], 22, "D6E2F0"),
            Texto(0.8, 4.6, 11.7, 0.5, ["Linguagens Formais e Autômatos · Trabalho do 1º Bimestre · CESUPA"], 16, "D6E2F0"),
            Texto(0.8, 5.1, 11.7, 0.5, ["Augusto Rodrigues  ·  Caue Jadão  ·  César Augusto"], 18, "FFFFFF"),
            Texto(0.8, 6.4, 11.7, 0.4, ["github.com/cauebarroso/Auditor-de-sensores-industriais-com-express-es-regulares"], 13, "AFC3DB"),
        ], "Abertura (todos). Apresentar a equipe e o tema em uma frase: um firewall que só deixa passar pacotes "
           "de sensores e atuadores que pertencem à linguagem de uma das cinco ERs.", capa=True),

        Slide("O problema e a solução", [
            Texto(0.5, 1.1, 6.6, 4.8, [
                "• Sensores e atuadores conversam com o supervisório (SCADA/IIoT) por pacotes de texto.",
                "• Um único símbolo errado vira uma leitura falsa ou uma ordem perigosa.",
                "• Solução: um firewall léxico. O pacote só passa se pertencer à linguagem de uma das cinco ERs.",
                "• Pacotes bloqueados recebem a causa provável e a coluna exata do erro, obtida simulando o AFNε.",
            ], 19),
            Texto(7.4, 1.1, 5.45, 0.4, ["Pacotes que chegam à rede"], 14, AZUL),
            Texto(7.4, 1.5, 5.45, 4.4, [
                "SEN-TMP-0101:TEMP=78.4C        ✔",
                "CMD ATU-VLV-0201 AJUSTAR 40%   ✔",
                "10.20.0.0/16                   ✔",
                "2026-09-30T06:42:10 WARN SEN-TMP-0101: ...",
                "",
                "SEN-TMP-0101:TEMP=1O4.7C       ✘",
                "     letra O no lugar do zero",
                "CMD ATU-VLV-0201 AJUSTAR 120%  ✘",
                "     ajuste acima de 100%",
                "2026-09-31T07:10:00 INFO ...   ✘",
                "     setembro tem 30 dias",
            ], 13, codigo=True, fundo=AZUL_CLARO),
        ], "Integrante 1 (~1 min). Contextualizar: redes industriais trocam texto; um caractere errado pode virar "
           "leitura falsa ou comando perigoso. Mostrar à direita exemplos reais dos dados de demonstração."),

        Slide("Entradas → processamento → saídas", [
            Texto(0.5, 1.2, 3.9, 3.3, ["**Entradas", "• Pacote digitado (modo interativo)", "• Pacote por argumento (--pacote)",
                                      "• Arquivo .txt em lote, com comentários #"], 16, fundo=AZUL_CLARO),
            Texto(4.7, 1.2, 3.9, 3.3, ["**Processamento", "• Normaliza (remove espaços das pontas)", "• re.fullmatch nas 5 ERs",
                                      "• Extrai campos dos válidos", "• Diagnostica os inválidos pelo AFNε"], 16, fundo=AZUL_CLARO),
            Texto(8.9, 1.2, 3.95, 3.3, ["**Saídas", "• [OK] / [ERRO] com coluna do erro", "• Relatório de conformidade",
                                       "• Filtros por categoria", "• Exportação JSON, CSV ou TXT"], 16, fundo=AZUL_CLARO),
            Tabela(0.5, 4.8, 12.35, [
                ["Módulo", "Responsabilidade"],
                ["expressoes.py", "As 5 ERs e as fichas (alfabeto, linguagem, ER formal)"],
                ["afn.py · construtor_afn.py", "AFNε: fecho-ε, simulação, construção de Thompson, prova de equivalência"],
                ["diagnostico.py · relatorio.py · auditor.py", "Diagnóstico, relatório/exportação e interface (menu e linha de comando)"],
            ], [4.4, 7.95], 13, codigo_colunas=(0,)),
        ], "Integrante 1 (~1 min). Requisito do trabalho: entradas, processamento e saídas. Destacar o tratamento de "
           "entradas vazias e inválidas (arquivo inexistente, vazio, binário, pasta) e a organização em módulos."),

        Slide("Notação formal × sintaxe no código", [
            Tabela(0.5, 1.15, 12.35, [
                ["Notação formal", "No Python", "Significado"],
                ["r s", "rs", "Concatenação"],
                ["r | s", "r|s", "União"],
                ["r*", "r*", "Fecho de Kleene (zero ou mais)"],
                ["r r*", "r+", "Fecho positivo (uma ou mais)"],
                ["(r | ε)", "r?", "Opcionalidade"],
                ["r r r r", "r{4}", "Repetição exata"],
                ["r r (r | ε)", "r{2,3}", "Repetição limitada = rr | rrr"],
                ["[0-9]", "[0-9]", "Classe finita: (0 | 1 | … | 9)"],
                ["⎵   e   .", "' '   e   \\.", "Espaço e ponto literais (o ponto do Python precisa de escape)"],
                ["⟨NOME⟩ = r", "—", "Definição regular: só um nome para a subexpressão"],
            ], [3.2, 2.6, 6.55], 15, 0.42, codigo_colunas=(0, 1)),
            Texto(0.5, 6.0, 12.35, 0.9, [
                "Sempre re.fullmatch (dispensa ^ e $). Proibidos e rejeitados pelo nosso analisador: "
                "\\d \\w \\s, ponto curinga, classes negadas, retroreferências e lookarounds."], 15, CINZA),
        ], "Integrante 1 (~40 s). Explicar que a ER formal e o código são duas escritas da mesma linguagem; os atalhos "
           "do Python (+, ?, {m,n}) são derivados dos operadores básicos. O próprio programa rejeita recursos não regulares."),

        slide_er("ER-01", "Identificador de um sensor (SEN) ou atuador (ATU); bloco reutilizado nas ERs 02, 03 e 05.",
                 DIAGRAMAS / "ER-01.png",
                 "Integrante 1 (~1 min). Mostrar a união SEN | ATU (dois ramos com ε), a terceira letra opcional "
                 "(desvio ε entre q12 e q13) e os 4 dígitos. L₁ é finita: 365.040.000 cadeias. É o bloco reutilizado nas ERs 2, 3 e 5.",
                 0.7),
        slide_er("ER-02", "Leitura de um sensor: temperatura (C), umidade (%) ou pressão (hPa), dentro da faixa.", None,
                 "Integrante 1 (~40 s). Três ramos na união, um por grandeza. Destacar ([1-9] | ε)[0-9], que impede "
                 "zero à esquerda, e o escape \\. da casa decimal (sem ele o ponto seria qualquer caractere).", 1.75),
        Slide("ER-02 — AFNε: o ID do sensor e a união das três medições", [
            Imagem(0.5, 1.05, 12.3, 1.3, DIAGRAMAS / "partes" / "ER-02_p1.png"),
            Imagem(0.5, 2.4, 12.3, 4.55, DIAGRAMAS / "partes" / "ER-02_p2.png"),
        ], "Integrante 1 (~40 s). Parte 1: o ID do sensor. Parte 2: três movimentos ε abrem os ramos TEMP, UMID "
           "e PRES. No ramo TEMP, o desvio ε paralelo ao '-' é o sinal opcional (- | ε)."),
        slide_er("ER-03", "Ordem para um atuador: ligar, desligar, abrir, fechar ou ajustar de 0 a 100%.",
                 DIAGRAMAS / "partes" / "ER-03_p2.png",
                 "Integrante 2 (~1 min). Cinco ações por união; o espaço é símbolo do alfabeto (⎵), por isso não usamos \\s. "
                 "⟨PCT⟩ = 100 | ([1-9] | ε)[0-9] cobre 0 a 100 sem zero à esquerda.", 0.95),
        slide_er("ER-04", "Endereço IPv4 com octetos de 0 a 255, sem zero à esquerda, e máscara CIDR opcional (/0 a /32).",
                 DIAGRAMAS / "partes" / "ER-04_p1.png",
                 "Integrante 2 (~1 min). O octeto é dividido em faixas disjuntas (250–255, 200–249, 100–199, 0–99). "
                 "No diagrama, cada ⟨OCT⟩ abre 4 ramos por ε. A máscara CIDR é opcional: (/⟨MASC⟩ | ε).", 0.95),
        slide_er("ER-05", "Registro de evento: data válida, hora, severidade, dispositivo e mensagem.", None,
                 "Integrante 3 (~40 s). Apresentar as definições regulares: ⟨DATA⟩, ⟨HORA⟩, ⟨SEV⟩, ⟨ID⟩ (reuso da ER-01) "
                 "e ⟨MSG⟩. É a única ER com fecho de Kleene.", 2.25),
        Slide("ER-05 — AFNε da data e da hora", [
            Imagem(0.5, 1.0, 12.3, 3.45, DIAGRAMAS / "partes" / "ER-05_p1.png"),
            Imagem(0.5, 4.5, 12.3, 2.45, DIAGRAMAS / "partes" / "ER-05_p2.png"),
        ], "Integrante 3 (~1 min). A data aceita o dia certo de cada mês: três casos na união (01–29 em qualquer mês, "
           "30 exceto fevereiro, 31 só nos meses de 31 dias). Validar dia por mês É regular, porque o conjunto de "
           "datas é finito. A parte 2 mostra a hora (00–23), os minutos/segundos (00–59) e a severidade."),

        Slide("ER-05 — a mensagem usa fecho de Kleene", [
            Texto(0.5, 1.05, 12.3, 0.9, [
                "⟨PAL⟩ = [A-Za-z0-9_.-] [A-Za-z0-9_.-]*        ⟨MSG⟩ = ⟨PAL⟩ (⎵ ⟨PAL⟩)*",
                'No código:  [A-Za-z0-9_.-]+( [A-Za-z0-9_.-]+)*'], 15, codigo=True, fundo=AZUL_CLARO),
            Imagem(0.5, 2.1, 12.3, 3.3, DIAGRAMAS / "partes" / "ER-05_p4.png"),
            Texto(0.5, 5.55, 12.3, 1.5, [
                "• Laço interno: o fecho de [A-Za-z0-9_.-]* (q104 ↔ q106). Laço externo: (⎵⟨PAL⟩)*, que volta de q112 para q107.",
                "• Por causa do fecho, L₅ é a única linguagem INFINITA do projeto; as outras quatro são finitas.",
                "• Rejeita mensagem vazia, dois espaços seguidos e símbolos fora de Σ₅ (ex.: pressão, com ã).",
            ], 16),
        ], "Integrante 3 (~1 min). Mostrar os dois laços do fecho de Kleene no AFNε. Explicar por que a ER-05 é infinita "
           "(há um ciclo que lê símbolos em um caminho até o estado final) e as outras são finitas."),

        Slide("Como garantimos que é a mesma linguagem", [
            Texto(0.5, 1.1, 6.4, 5.6, [
                "**1. Código → AFNε (Thompson)",
                "O AFNε é construído a partir do padrão exato do código.",
                "**2. ER formal ≡ código: prova",
                "Busca em largura sobre o produto dos dois AFNε (construção de subconjuntos). Nenhuma cadeia distingue as linguagens.",
                "**3. Arquivo .jff ≡ código: prova",
                "O .jff é lido de volta e comparado pelo mesmo algoritmo.",
                "**4. Executado no próprio JFLAP 7.1",
                "80 de 80 cadeias de teste com o resultado esperado.",
            ], 16),
            Tabela(7.2, 1.2, 5.65, [
                ["Verificação", "Resultado"],
                ["ER formal ≡ código", "5/5 provadas"],
                [".jff ≡ código", "5/5 provadas"],
                ["AFNε × re (aleatórias)", "10.000 cadeias, 0 divergências"],
                ["JFLAP 7.1", "80/80"],
                ["Linguagens disjuntas", "interseções vazias"],
            ], [2.9, 2.75], 14, 0.45),
            Texto(7.2, 4.3, 5.65, 1.3, [
                "Achado: o JFLAP 7.1 trata rótulos com '[' como intervalo ([a-z]) e falha com o símbolo '['. "
                "Por isso a severidade é INFO, não [INFO]: o .jff reconhece a linguagem exata, sem símbolos substitutos."],
                13, CINZA, fundo=AZUL_CLARO),
        ], "Integrante 2 (~1 min 30 s). Esta é a regra central do guia. Explicar que não comparamos só alguns exemplos: "
           "o algoritmo percorre todos os pares de estados alcançáveis, então é uma prova de equivalência."),

        Slide("Demonstração", [
            Texto(0.5, 1.05, 12.35, 1.45, [
                "python src/auditor.py dados/turno_caldeira.txt          # lote + relatório + filtros",
                'python src/auditor.py --pacote "CMD ATU-VLV-0201 AJUSTAR 120%"',
                'python src/auditor.py --explicar "ATU-VLV-0420"            # simulação do AFNε',
                "JFLAP: automatos/jflap/ER-03.jff → Input › Multiple Run → Load Inputs (entradas/ER-03.txt)",
            ], 14, codigo=True, fundo=AZUL_CLARO),
            Texto(0.5, 2.75, 12.35, 4.2, saida_demonstracao(), 13, codigo=True),
        ], "Integrante 3 (~2 min). Rodar ao vivo: lote com o turno da caldeira, um pacote rejeitado com o apontador "
           "de coluna, a simulação passo a passo e um Multiple Run no JFLAP. Se o professor pedir uma cadeia nova, "
           "usar --pacote ou o modo interativo (opção 2)."),

        slide_testes(),

        Slide("Limitações e possíveis melhorias", [
            Texto(0.5, 1.1, 6.1, 5.6, [
                "**Limitações (falsos resultados)",
                "• A validação é léxica: um ID bem formado passa mesmo se o dispositivo não existir.",
                "• −0C é aceito; uma leitura real acima de 199.9 °C é bloqueada.",
                "• 29/02 é aceito em qualquer ano.",
                "• A mensagem do log não aceita acentos.",
            ], 20),
            Texto(6.85, 1.1, 6.0, 5.6, [
                "**Melhorias",
                "• Ano bissexto: também é regular (anos de 4 dígitos formam um conjunto finito), mas o AFNε cresceria muito.",
                "• Faixas configuráveis por tipo de sensor.",
                "• Leitura contínua de porta serial ou socket.",
                "• Minimizar o AFD (Hopcroft) e comparar com o AFNε de Thompson.",
            ], 20),
        ], "Integrante 3 (~40 s). Separar limitação léxica de limitação semântica: tudo o que a ER promete, ela cumpre; "
           "o que depende do significado do dado fica fora."),

        Slide("Contribuições e uso de IA", [
            Tabela(0.5, 1.15, 12.35, [
                ["Integrante", "Contribuições"],
                ["Augusto Rodrigues", "ER-01 e ER-02 (fichas, AFNε, testes); diagnóstico de falhas"],
                ["Caue Jadão", "ER-03 e ER-04; revisão geral, verificação de equivalência, relatório e repositório"],
                ["César Augusto", "ER-05; modo lote, relatório de conformidade, exportação e apresentação"],
            ], [3.2, 9.15], 15, 0.5),
            Texto(0.5, 3.6, 12.35, 3.0, [
                "**Uso de IA (declarado no relatório e no README)",
                "• Antigravity: versão inicial do código, dos testes e dos primeiros .jff.",
                "• Claude (Anthropic): revisão das ERs e da notação formal, construção de Thompson, verificação de "
                "equivalência, testes, relatório e slides.",
                "• As decisões de projeto são da equipe; todos compreendem, explicam e modificam o conteúdo.",
            ], 17),
        ], "Todos (~30 s). Cada integrante diz em uma frase o que fez."),

        Slide("", [
            Texto(0.8, 2.6, 11.7, 1.0, ["Obrigado!"], 44, "FFFFFF"),
            Texto(0.8, 3.7, 11.7, 0.6, ["Perguntas?"], 26, "D6E2F0"),
            Texto(0.8, 6.4, 11.7, 0.4, ["github.com/cauebarroso/Auditor-de-sensores-industriais-com-express-es-regulares"], 13, "AFC3DB"),
        ], "Encerramento. Deixar o terminal aberto no modo interativo para testar cadeias pedidas pelo professor.", capa=True),
    ]


# ---------------------------------------------------------------------- #
# Renderização PPTX
# ---------------------------------------------------------------------- #
def _cor(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


def _texto_pptx(slide, el: Texto) -> None:
    caixa = slide.shapes.add_textbox(Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h))
    if el.fundo:
        caixa.fill.solid()
        caixa.fill.fore_color.rgb = _cor(el.fundo)
    quadro = caixa.text_frame
    quadro.word_wrap = True
    quadro.vertical_anchor = MSO_ANCHOR.TOP
    margem = Inches(0.12 if el.fundo else 0.04)
    quadro.margin_left = quadro.margin_right = margem
    quadro.margin_top = quadro.margin_bottom = Inches(0.08 if el.fundo else 0.02)
    for i, linha in enumerate(el.linhas):
        p = quadro.paragraphs[0] if i == 0 else quadro.add_paragraph()
        negrito = linha.startswith("**")
        p.text = linha[2:] if negrito else linha
        p.alignment = {"esquerda": PP_ALIGN.LEFT, "direita": PP_ALIGN.RIGHT, "centro": PP_ALIGN.CENTER}[el.alinhamento]
        if negrito and i > 0:
            p.space_before = Pt(10)
        elif not el.codigo:
            p.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(el.tamanho + (1 if negrito else 0))
            run.font.bold = negrito
            run.font.name = FONTE_CODIGO if el.codigo else FONTE
            run.font.color.rgb = _cor(AZUL if negrito else el.cor)


def _imagem_pptx(slide, el: Imagem) -> None:
    largura, altura = Image.open(el.caminho).size
    escala = min(el.w / largura, el.h / altura)
    w, h = largura * escala, altura * escala
    slide.shapes.add_picture(str(el.caminho), Inches(el.x + (el.w - w) / 2), Inches(el.y + (el.h - h) / 2),
                             Inches(w), Inches(h))


def _tabela_pptx(slide, el: Tabela) -> None:
    forma = slide.shapes.add_table(len(el.linhas), len(el.larguras), Inches(el.x), Inches(el.y),
                                   Inches(sum(el.larguras)), Inches(el.altura_linha * len(el.linhas)))
    tabela = forma.table
    for j, largura in enumerate(el.larguras):
        tabela.columns[j].width = Inches(largura)
    for i, linha in enumerate(el.linhas):
        tabela.rows[i].height = Inches(el.altura_linha)
        for j, valor in enumerate(linha):
            celula = tabela.cell(i, j)
            celula.text = valor
            celula.margin_left = celula.margin_right = Inches(0.08)
            celula.margin_top = celula.margin_bottom = Inches(0.03)
            celula.vertical_anchor = MSO_ANCHOR.MIDDLE
            celula.fill.solid()
            celula.fill.fore_color.rgb = _cor(AZUL if i == 0 else ("F7F9FC" if i % 2 else "FFFFFF"))
            for run in celula.text_frame.paragraphs[0].runs:
                run.font.size = Pt(el.tamanho)
                run.font.bold = i == 0
                run.font.name = FONTE_CODIGO if (i > 0 and j in el.codigo_colunas) else FONTE
                run.font.color.rgb = _cor("FFFFFF" if i == 0 else "1B1F24")


def gerar_pptx(lista: list[Slide], destino: Path) -> None:
    apresentacao = Presentation()
    apresentacao.slide_width, apresentacao.slide_height = Inches(LARGURA), Inches(ALTURA)
    vazio = apresentacao.slide_layouts[6]
    for numero, s in enumerate(lista, start=1):
        slide = apresentacao.slides.add_slide(vazio)
        fundo = slide.background.fill
        fundo.solid()
        fundo.fore_color.rgb = _cor(AZUL if s.capa else "FFFFFF")
        if not s.capa:
            faixa = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(LARGURA), Inches(0.85))
            faixa.fill.solid()
            faixa.fill.fore_color.rgb = _cor(AZUL)
            faixa.line.fill.background()
            _texto_pptx(slide, Texto(0.45, 0.14, 10.0, 0.6, [s.titulo], 26, "FFFFFF"))
            _texto_pptx(slide, Texto(11.9, 7.05, 1.0, 0.3, [f"{numero}/{len(lista)}"], 11, CINZA, "direita"))
        for el in s.elementos:
            {Texto: _texto_pptx, Imagem: _imagem_pptx, Tabela: _tabela_pptx}[type(el)](slide, el)
        if s.notas:
            slide.notes_slide.notes_text_frame.text = s.notas
    apresentacao.save(destino)


# ---------------------------------------------------------------------- #
# Renderização HTML -> PDF
# ---------------------------------------------------------------------- #
def _caixa(el, conteudo: str, extra: str = "") -> str:
    return (f'<div style="position:absolute;left:{el.x}in;top:{el.y}in;width:{el.w}in;'
            f'{f"height:{el.h}in;" if hasattr(el, "h") else ""}{extra}">{conteudo}</div>')


def _texto_html(el: Texto) -> str:
    paragrafos = []
    for i, linha in enumerate(el.linhas):
        negrito = linha.startswith("**")
        texto = html.escape(linha[2:] if negrito else linha) or "&nbsp;"
        estilo = (f"font-weight:700;color:#{AZUL};font-size:{el.tamanho + 1}pt;"
                  f"{'margin-top:10pt;' if i else ''}") if negrito else ""
        paragrafos.append(f'<p style="margin:0 0 {0 if el.codigo else 4}pt;{estilo}">{texto}</p>')
    fonte = f"'{FONTE_CODIGO}',monospace;white-space:pre-wrap" if el.codigo else f"'{FONTE}',sans-serif"
    alinhamento = {"esquerda": "left", "direita": "right", "centro": "center"}[el.alinhamento]
    preenchimento = f"background:#{el.fundo};padding:0.08in 0.12in;" if el.fundo else "padding:0.02in 0.04in;"
    return _caixa(el, "".join(paragrafos),
                  f"box-sizing:border-box;{preenchimento}font-family:{fonte};font-size:{el.tamanho}pt;"
                  f"color:#{el.cor};text-align:{alinhamento};line-height:1.2;overflow:hidden;")


def _imagem_html(el: Imagem) -> str:
    return _caixa(el, f'<img src="{el.caminho.resolve().as_uri()}" '
                      'style="width:100%;height:100%;object-fit:contain">')


def _tabela_html(el: Tabela) -> str:
    linhas = []
    for i, linha in enumerate(el.linhas):
        celulas = []
        for j, valor in enumerate(linha):
            fonte = FONTE_CODIGO if (i > 0 and j in el.codigo_colunas) else FONTE
            fundo = AZUL if i == 0 else ("F7F9FC" if i % 2 else "FFFFFF")
            celulas.append(
                f'<td style="width:{el.larguras[j]}in;height:{el.altura_linha}in;background:#{fundo};'
                f"font-family:'{fonte}';font-weight:{700 if i == 0 else 400};color:#{'FFFFFF' if i == 0 else '1B1F24'};"
                f'padding:0.03in 0.08in;border:1px solid #C9D4E2;box-sizing:border-box">{html.escape(valor)}</td>')
        linhas.append(f"<tr>{''.join(celulas)}</tr>")
    return _caixa(el, f'<table style="border-collapse:collapse;font-size:{el.tamanho}pt">{"".join(linhas)}</table>')


def gerar_pdf(lista: list[Slide], destino: Path) -> Path | None:
    paginas = []
    for numero, s in enumerate(lista, start=1):
        partes = []
        if not s.capa:
            partes.append(f'<div style="position:absolute;left:0;top:0;width:{LARGURA}in;height:0.85in;background:#{AZUL}"></div>')
            partes.append(_texto_html(Texto(0.45, 0.14, 10.0, 0.6, [s.titulo], 26, "FFFFFF")))
            partes.append(_texto_html(Texto(11.9, 7.05, 1.0, 0.3, [f"{numero}/{len(lista)}"], 11, CINZA, "direita")))
        for el in s.elementos:
            partes.append({Texto: _texto_html, Imagem: _imagem_html, Tabela: _tabela_html}[type(el)](el))
        fundo = AZUL if s.capa else "FFFFFF"
        paginas.append(f'<section style="position:relative;width:{LARGURA}in;height:{ALTURA}in;'
                       f'background:#{fundo};overflow:hidden;page-break-after:always">{"".join(partes)}</section>')
    documento = (f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><style>'
                 f"@page{{size:{LARGURA}in {ALTURA}in;margin:0}}body{{margin:0}}"
                 f"</style></head><body>{''.join(paginas)}</body></html>")
    return html_para_pdf(documento, destino)


def main() -> None:
    lista = slides()
    destino = RAIZ / "docs"
    destino.mkdir(exist_ok=True)
    gerar_pptx(lista, destino / "Apresentacao.pptx")
    print("Gerado: docs/Apresentacao.pptx")
    pdf = gerar_pdf(lista, destino / "Apresentacao.pdf")
    print("Gerado: docs/Apresentacao.pdf" if pdf else "[AVISO] Edge/Chrome não encontrado: PDF não gerado.")


if __name__ == "__main__":
    main()
